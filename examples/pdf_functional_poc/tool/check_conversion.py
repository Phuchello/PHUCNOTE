"""Optional desktop inspection; requires existing pymupdf/python-pptx/Pillow."""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re
import xml.etree.ElementTree as ET
import pymupdf
from pptx import Presentation
from PIL import Image, ImageOps

repo = Path(__file__).resolve().parents[3]
fixtures = repo / 'fixtures/document_lab'
source = fixtures / 'lecture_slides_network_security.pptx'
converted = fixtures / 'lecture_slides_network_security_converted.pdf'
deck = Presentation(source)
pdf = pymupdf.open(converted)
normalize = lambda text: re.sub(r'\s+', '', text)
data = {'converter': 'Microsoft PowerPoint 16.0, SaveAs ppSaveAsPDF=32, source read-only',
        'sourceSHA256': hashlib.sha256(source.read_bytes()).hexdigest(),
        'outputSHA256': hashlib.sha256(converted.read_bytes()).hexdigest(),
        'outputBytes': converted.stat().st_size, 'slides': [],
        'embeddedPDFFonts': sorted({font[3] for page in pdf for font in page.get_fonts()})}
with ZipFile(source) as archive:
    theme = ET.fromstring(archive.read('ppt/theme/theme1.xml'))
    data['themeLatinFonts'] = [node.attrib['typeface'] for node in theme.iter()
                               if node.tag.endswith('}latin')]
sheet = Image.new('RGB', (1440, 1290), '#eee8da')
for index, slide in enumerate(deck.slides):
    page = pdf[index]
    texts = [paragraph.text for shape in slide.shapes if shape.has_text_frame
             for paragraph in shape.text_frame.paragraphs if paragraph.text.strip()]
    for shape in slide.shapes:
        if shape.has_table:
            texts += [cell.text for row in shape.table.rows for cell in row.cells if cell.text.strip()]
    rendered = normalize(page.get_text())
    missing = [text for text in texts if normalize(text) not in rendered]
    data['slides'].append({'page': index+1, 'widthPt': page.rect.width,
        'heightPt': page.rect.height, 'checkedTextItems': len(texts),
        'missingTextItems': missing, 'tables': sum(shape.has_table for shape in slide.shapes)})
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(.75,.75))
    image = Image.frombytes('RGB', (pixmap.width,pixmap.height), pixmap.samples)
    sheet.paste(ImageOps.contain(image,(700,394)), ((index%2)*720,(index//2)*430))
output = repo / 'artifacts/pdf-functional-poc'
output.mkdir(parents=True, exist_ok=True)
sheet.save(output/'conversion-contact.png', optimize=True)
(output/'conversion-trace.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(data,ensure_ascii=False,indent=2))
